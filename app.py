"""
Ứng dụng GenAI Content Rewriter & Text-to-Speech
Tác giả: Senior Python Developer
Framework: Streamlit + Google Gemini + Edge TTS
"""

import streamlit as st
import google.generativeai as genai
from bs4 import BeautifulSoup
import requests
from newspaper import Article
import edge_tts
import asyncio
import os
from io import BytesIO
from docx import Document
from docx.shared import Pt
from datetime import datetime
import re

# ==================== CẤU HÌNH TRANG ====================
st.set_page_config(
    page_title="GenAI Content Rewriter",
    page_icon="🤖",
    layout="wide",
    initial_sidebar_state="expanded"
)

# ==================== CSS TÙY CHỈNH ====================
st.markdown("""
<style>
    .main-header {
        font-size: 2.5rem;
        color: #1E88E5;
        text-align: center;
        margin-bottom: 2rem;
    }
    .stButton>button {
        width: 100%;
        background-color: #1E88E5;
        color: white;
        font-weight: bold;
    }
    .success-box {
        padding: 1rem;
        background-color: #D4EDDA;
        border-radius: 5px;
        margin: 1rem 0;
    }
</style>
""", unsafe_allow_html=True)

# ==================== HÀM SCRAPE NỘI DUNG ====================
def scrape_content(url):
    """
    Scrape nội dung từ URL sử dụng newspaper3k và BeautifulSoup fallback
    """
    try:
        # Thử với newspaper3k trước (tốt hơn cho bài báo)
        article = Article(url)
        article.download()
        article.parse()
        
        if article.text and len(article.text) > 100:
            return {
                'title': article.title or 'Không có tiêu đề',
                'content': article.text,
                'method': 'newspaper3k'
            }
    except Exception as e:
        st.warning(f"Newspaper3k failed: {e}. Trying BeautifulSoup...")
    
    try:
        # Fallback với BeautifulSoup
        headers = {
            'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'
        }
        response = requests.get(url, headers=headers, timeout=10)
        response.raise_for_status()
        
        soup = BeautifulSoup(response.content, 'html.parser')
        
        # Loại bỏ script và style
        for script in soup(['script', 'style', 'nav', 'footer', 'header']):
            script.decompose()
        
        # Lấy text từ paragraphs
        paragraphs = soup.find_all('p')
        content = '\n\n'.join([p.get_text().strip() for p in paragraphs if p.get_text().strip()])
        
        title = soup.find('h1')
        title_text = title.get_text().strip() if title else 'Không có tiêu đề'
        
        if content and len(content) > 100:
            return {
                'title': title_text,
                'content': content,
                'method': 'BeautifulSoup'
            }
        else:
            raise ValueError("Không thể trích xuất đủ nội dung từ URL")
            
    except Exception as e:
        raise Exception(f"Lỗi khi scrape URL: {str(e)}")

# ==================== HÀM CHIA NỘI DUNG THÀNH CHUNKS ====================
def chunk_text(text, max_chars=30000):
    """
    Chia text thành các chunks nhỏ hơn để tránh vượt quá token limit
    """
    # Chia theo đoạn văn trước
    paragraphs = text.split('\n\n')
    chunks = []
    current_chunk = ""
    
    for para in paragraphs:
        if len(current_chunk) + len(para) < max_chars:
            current_chunk += para + "\n\n"
        else:
            if current_chunk:
                chunks.append(current_chunk.strip())
            current_chunk = para + "\n\n"
    
    if current_chunk:
        chunks.append(current_chunk.strip())
    
    return chunks

# ==================== HÀM XỬ LÝ GEMINI ====================
def rewrite_with_gemini(content, style, target_length, api_key):
    """
    Viết lại nội dung sử dụng Google Gemini API
    """
    try:
        genai.configure(api_key=api_key)
        model = genai.GenerativeModel('gemini-pro')
        
        # Định nghĩa các phong cách
        style_prompts = {
            'Hài hước': 'với phong cách hài hước, dí dỏm, thu hút người đọc',
            'Chuyên gia': 'với phong cách chuyên nghiệp, chuyên gia, mang tính học thuật',
            'Kể chuyện': 'với phong cách kể chuyện, sinh động, có cảm xúc',
            'Tối giản': 'với phong cách tối giản, súc tích, đi thẳng vào vấn đề',
            'Học thuật': 'với phong cách học thuật, nghiêm túc, có trích dẫn'
        }
        
        style_text = style_prompts.get(style, style_prompts['Chuyên gia'])
        
        # Kiểm tra độ dài nội dung
        word_count = len(content.split())
        
        if word_count > 20000:
            # Chia nhỏ nội dung
            chunks = chunk_text(content, max_chars=25000)
            st.info(f"Nội dung dài, đang xử lý {len(chunks)} phần...")
            
            rewritten_parts = []
            progress_bar = st.progress(0)
            
            for i, chunk in enumerate(chunks):
                prompt = f"""
Bạn là một chuyên gia viết lại nội dung. Hãy viết lại nội dung sau đây {style_text}.

QUAN TRỌNG: Đây là phần {i+1}/{len(chunks)} của bài viết. Giữ tính liên kết với các phần trước/sau.

Yêu cầu:
- Viết bằng tiếng Việt
- Độ dài: khoảng {target_length // len(chunks)} từ cho phần này
- Giữ nguyên ý chính và thông tin quan trọng
- Tối ưu hóa cấu trúc và cách diễn đạt
- Không thêm thông tin không có trong bài gốc

Nội dung gốc:
{chunk}

Hãy bắt đầu viết lại:
"""
                
                response = model.generate_content(prompt)
                rewritten_parts.append(response.text)
                progress_bar.progress((i + 1) / len(chunks))
            
            return '\n\n'.join(rewritten_parts)
        
        else:
            # Xử lý bình thường nếu nội dung ngắn
            prompt = f"""
Bạn là một chuyên gia viết lại nội dung. Hãy viết lại nội dung sau đây {style_text}.

Yêu cầu:
- Viết bằng tiếng Việt
- Độ dài mục tiêu: khoảng {target_length} từ
- Giữ nguyên ý chính và thông tin quan trọng
- Tối ưu hóa cấu trúc và cách diễn đạt
- Không thêm thông tin không có trong bài gốc

Nội dung gốc:
{content}

Hãy bắt đầu viết lại:
"""
            
            response = model.generate_content(prompt)
            return response.text
            
    except Exception as e:
        raise Exception(f"Lỗi khi xử lý với Gemini: {str(e)}")

# ==================== HÀM TEXT-TO-SPEECH ====================
async def text_to_speech_async(text, voice, output_file):
    """
    Chuyển đổi text thành speech sử dụng Edge TTS
    """
    communicate = edge_tts.Communicate(text, voice)
    await communicate.save(output_file)

def text_to_speech(text, voice='vi-VN-HoaiMyNeural'):
    """
    Wrapper đồng bộ cho text_to_speech_async
    """
    try:
        # Chia text thành các phần nhỏ nếu quá dài (Edge TTS có giới hạn)
        max_chars = 5000
        if len(text) > max_chars:
            chunks = [text[i:i+max_chars] for i in range(0, len(text), max_chars)]
            audio_files = []
            
            for i, chunk in enumerate(chunks):
                output_file = f"temp_audio_part_{i}.mp3"
                asyncio.run(text_to_speech_async(chunk, voice, output_file))
                audio_files.append(output_file)
            
            return audio_files
        else:
            output_file = "temp_audio.mp3"
            asyncio.run(text_to_speech_async(text, voice, output_file))
            return [output_file]
            
    except Exception as e:
        raise Exception(f"Lỗi khi tạo audio: {str(e)}")

# ==================== HÀM TẠO FILE DOCX ====================
def create_docx(title, content):
    """
    Tạo file DOCX từ nội dung
    """
    doc = Document()
    
    # Thêm tiêu đề
    heading = doc.add_heading(title, 0)
    heading.alignment = 1  # Center
    
    # Thêm thông tin
    doc.add_paragraph(f"Được tạo bởi GenAI Content Rewriter")
    doc.add_paragraph(f"Ngày: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    doc.add_paragraph("")
    
    # Thêm nội dung
    for paragraph in content.split('\n\n'):
        if paragraph.strip():
            p = doc.add_paragraph(paragraph.strip())
            p.style.font.size = Pt(12)
    
    # Lưu vào BytesIO
    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio

# ==================== GIAO DIỆN CHÍNH ====================
def main():
    st.markdown('<h1 class="main-header">🤖 GenAI Content Rewriter & TTS</h1>', unsafe_allow_html=True)
    
    # ========== SIDEBAR CẤU HÌNH ==========
    with st.sidebar:
        st.header("⚙️ Cấu hình")
        
        # API Key
        api_key = st.text_input(
            "Google Gemini API Key",
            type="password",
            help="Nhập API key từ https://makersuite.google.com/app/apikey"
        )
        
        if not api_key:
            st.warning("⚠️ Vui lòng nhập API Key để sử dụng")
        
        st.divider()
        
        # Cấu hình phong cách
        st.subheader("📝 Phong cách viết")
        style = st.selectbox(
            "Chọn phong cách",
            ['Hài hước', 'Chuyên gia', 'Kể chuyện', 'Tối giản', 'Học thuật']
        )
        
        # Cấu hình độ dài
        st.subheader("📏 Độ dài mục tiêu")
        target_length = st.slider(
            "Số từ",
            min_value=1000,
            max_value=50000,
            value=5000,
            step=500
        )
        
        # Cấu hình giọng đọc
        st.subheader("🎙️ Giọng đọc")
        voice_option = st.radio(
            "Chọn giọng",
            ['Nam (Nam Minh)', 'Nữ (Hoài My)']
        )
        
        voice_map = {
            'Nam (Nam Minh)': 'vi-VN-NamMinhNeural',
            'Nữ (Hoài My)': 'vi-VN-HoaiMyNeural'
        }
        selected_voice = voice_map[voice_option]
        
        st.divider()
        st.info("💡 **Mẹo**: Nội dung dài sẽ tự động được chia nhỏ để xử lý tốt hơn")
    
    # ========== PHẦN CHÍNH ==========
    tab1, tab2, tab3 = st.tabs(["📥 Input", "✍️ Rewrite", "🔊 Audio"])
    
    with tab1:
        st.subheader("Nhập URL để bắt đầu")
        
        url = st.text_input(
            "URL của bài viết",
            placeholder="https://example.com/article",
            help="Nhập đường link bài viết bạn muốn viết lại"
        )
        
        col1, col2 = st.columns([1, 4])
        with col1:
            scrape_button = st.button("🌐 Scrape Content", use_container_width=True)
        
        if scrape_button and url:
            if not api_key:
                st.error("❌ Vui lòng nhập API Key trong Sidebar")
            else:
                with st.spinner("Đang scrape nội dung..."):
                    try:
                        result = scrape_content(url)
                        
                        # Lưu vào session state
                        st.session_state['scraped_data'] = result
                        st.session_state['url'] = url
                        
                        st.success(f"✅ Scrape thành công bằng {result['method']}")
                        
                        # Hiển thị preview
                        st.markdown("### 📄 Nội dung đã scrape")
                        st.markdown(f"**Tiêu đề:** {result['title']}")
                        st.markdown(f"**Số từ:** {len(result['content'].split())}")
                        
                        with st.expander("Xem nội dung gốc"):
                            st.text_area("", result['content'], height=300)
                        
                    except Exception as e:
                        st.error(f"❌ Lỗi: {str(e)}")
        
        elif scrape_button:
            st.warning("⚠️ Vui lòng nhập URL")
    
    with tab2:
        st.subheader("Viết lại nội dung với AI")
        
        if 'scraped_data' in st.session_state:
            col1, col2 = st.columns([1, 4])
            with col1:
                rewrite_button = st.button("✨ Rewrite Now", use_container_width=True)
            
            if rewrite_button:
                with st.spinner(f"Đang viết lại với phong cách {style}..."):
                    try:
                        rewritten = rewrite_with_gemini(
                            st.session_state['scraped_data']['content'],
                            style,
                            target_length,
                            api_key
                        )
                        
                        st.session_state['rewritten_content'] = rewritten
                        st.session_state['rewrite_style'] = style
                        
                        st.success("✅ Viết lại thành công!")
                        
                    except Exception as e:
                        st.error(f"❌ Lỗi: {str(e)}")
            
            # Hiển thị kết quả nếu có
            if 'rewritten_content' in st.session_state:
                st.markdown("### 📝 Nội dung đã viết lại")
                st.markdown(f"**Phong cách:** {st.session_state.get('rewrite_style', 'N/A')}")
                st.markdown(f"**Số từ:** {len(st.session_state['rewritten_content'].split())}")
                
                # Hiển thị với expander cho nội dung dài
                with st.expander("Xem toàn bộ nội dung", expanded=True):
                    st.markdown(st.session_state['rewritten_content'])
                
                # Nút download
                col1, col2, col3 = st.columns(3)
                
                with col1:
                    st.download_button(
                        label="📄 Tải về TXT",
                        data=st.session_state['rewritten_content'],
                        file_name=f"rewritten_{datetime.now().strftime('%Y%m%d_%H%M%S')}.txt",
                        mime="text/plain"
                    )
                
                with col2:
                    docx_file = create_docx(
                        st.session_state['scraped_data']['title'],
                        st.session_state['rewritten_content']
                    )
                    st.download_button(
                        label="📘 Tải về DOCX",
                        data=docx_file,
                        file_name=f"rewritten_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    )
        else:
            st.info("ℹ️ Vui lòng scrape nội dung ở tab Input trước")
    
    with tab3:
        st.subheader("Chuyển đổi thành Audio")
        
        if 'rewritten_content' in st.session_state:
            col1, col2 = st.columns([1, 4])
            with col1:
                tts_button = st.button("🎙️ Generate Audio", use_container_width=True)
            
            if tts_button:
                with st.spinner("Đang tạo audio..."):
                    try:
                        audio_files = text_to_speech(
                            st.session_state['rewritten_content'],
                            selected_voice
                        )
                        
                        st.session_state['audio_files'] = audio_files
                        st.success(f"✅ Tạo audio thành công! ({len(audio_files)} phần)")
                        
                    except Exception as e:
                        st.error(f"❌ Lỗi: {str(e)}")
            
            # Hiển thị audio player nếu có
            if 'audio_files' in st.session_state:
                st.markdown("### 🔊 Audio đã tạo")
                
                for i, audio_file in enumerate(st.session_state['audio_files']):
                    if os.path.exists(audio_file):
                        st.markdown(f"**Phần {i+1}/{len(st.session_state['audio_files'])}**")
                        
                        # Player
                        audio_bytes = open(audio_file, 'rb').read()
                        st.audio(audio_bytes, format='audio/mp3')
                        
                        # Download button
                        st.download_button(
                            label=f"💾 Tải về Part {i+1}",
                            data=audio_bytes,
                            file_name=f"audio_part_{i+1}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.mp3",
                            mime="audio/mp3",
                            key=f"download_{i}"
                        )
                        
                        st.divider()
        else:
            st.info("ℹ️ Vui lòng viết lại nội dung ở tab Rewrite trước")

# ==================== CHẠY ỨNG DỤNG ====================
if __name__ == "__main__":
    # Khởi tạo session state
    if 'scraped_data' not in st.session_state:
        st.session_state['scraped_data'] = None
    
    main()
